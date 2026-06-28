#!/usr/bin/env python3
"""Genera SocialPerks_Restaurants_Vienna.xlsx"""

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

RESTAURANTS = [
    ("Vogel Kaffee", "2. Bezirk", "Bruno-Marek-Allee 19, 1020 Wien", "hello@vogelkaffee.at", "+43 660 6307479", "vogelkaffee.at", "@vogelkaffee", "Specialty Coffee / Micro-Roastery", "Indipendente", "✅ Sì", "Micro-torrefazione speciality nel 2. distretto, molto instagrammabile"),
    ("Café Ansari", "2. Bezirk", "Praterstraße 15, 1020 Wien", "cafe@cafeansari.at", "+43 1 2765102", "cafeansari.at", "@cafeansari", "Café / Fusion", "Indipendente", "✅ Sì", "Café trendy bohémien nel 2. distretto, cucina creativa"),
    ("Cà Phê Lalot", "1. Bezirk", "Wipplingerstraße 25, 1010 Wien", "hello@caphelalot.com", "", "caphelalot.com", "@caphe.lalot", "Café / Asian Fusion", "Indipendente", "✅ Sì", "Caffè diurno con focus asiatico, piatti creativi, solo walk-in"),
    ("Crème de la Crème", "4./8. Bezirk", "Kettenbrückengasse 20, 1040 Wien", "hallo@cremedelacreme.at", "+43 660 2833769", "cremedelacreme.at", "@cremedelacremevienna", "Patisserie / Café", "Piccola catena", "⚠️ Medio", "Patisserie francese molto instagrammabile, pasticcini artigianali"),
    ("Café Vollpension", "4./1. Bezirk", "Schleifmühlgasse 16, 1040 Wien", "info@vollpension.wien", "", "vollpension.wien", "@vollpension_wien", "Café / Kuchen", "Piccola catena", "⚠️ Medio", "Generationencafé con torte fatte da nonni, concept unico e fotografabile"),
    ("Bali Brunch", "7. Bezirk", "Lerchenfelderstraße 1-3 (25hours Hotel), 1070 Wien", "balibrunch@gmail.com", "+43 677 64066513", "balibrunch.at", "@bali_brunch", "Brunch / Asian-Inspired", "Indipendente", "✅ Sì", "Brunch ispirato a Bali sulla rooftop del 25hours Hotel, 45K follower IG"),
    ("Café Drechsler", "6. Bezirk", "Linke Wienzeile 22, 1060 Wien", "office@drechsler-wien.at", "+43 676 5962730", "drechsler-wien.at", "@drechsler_wienzeile", "Café / Wiener Kaffeehaus", "Indipendente", "✅ Sì", "Storico caffè viennese moderno al Naschmarkt, interior design iconico"),
    ("Café Ulrich", "7. Bezirk", "St.-Ulrichs-Platz 1, 1070 Wien", "hallo@ulrichwien.at", "+43 1 9612782", "ulrichwien.at", "@ulricherichwien", "Café / Restaurant / Bar", "Indipendente", "✅ Sì", "Café e bar trendy su piazza pittoresca nel 7. distretto"),
    ("Erich Café Bar", "7. Bezirk", "Neustiftgasse 27, 1070 Wien", "hallo@erichwien.at", "+43 1 8906400", "erichwien.at", "@ulricherichwien", "Café / Bar", "Indipendente", "✅ Sì", "Fratello di Café Ulrich, stessa atmosfera trendy nel 7. distretto"),
    ("Café Menta", "3. Bezirk", "Radetzkyplatz 4, 1030 Wien", "reservierung@cafementa.at", "+43 1 9668423", "cafementa.at", "@cafe_menta", "Café / Restaurant", "Indipendente", "✅ Sì", "Café con design industriale al Radetzkyplatz, atmosfera accogliente"),
    ("Café Liebling", "7. Bezirk", "Zollergasse 6, 1070 Wien", "info@cafeliebling.at", "+43 1 9905877", "cafeliebling.at", "@cafeliebling", "Café / Brunch / Mediterranean", "Indipendente", "✅ Sì", "Café brunch mediterraneo nel cuore del 7. distretto"),
    ("Adlerhof", "7. Bezirk", "Burggasse 51, 1070 Wien", "post@adlerhof.wien", "+43 1 5224905", "adlerhof.wien", "@adlerhof_wien", "Café / Bistro / Bar", "Indipendente", "✅ Sì", "Bistro con scale a chiocciola e wallpaper vintage, uno dei più instagrammabili di Vienna"),
    ("Café Schopenhauer", "18. Bezirk", "Staudgasse 1, 1180 Wien", "office@cafeschopenhauer.at", "+43 1 4063288", "cafeschopenhauer.at", "@cafeschopenhauer", "Café / Libreria", "Indipendente", "✅ Sì", "Caffè storico rinnovato con libreria integrata, atmosfera letteraria instagrammabile"),
    ("Café Weimar", "9. Bezirk", "Währinger Straße 68, 1090 Wien", "service@cafeweimar.at", "+43 1 3171206", "cafeweimar.at", "@cafeweimar", "Café / Wiener Kaffeehaus", "Indipendente", "✅ Sì", "Storico caffè viennese nell'Alsergrund, atmosfera classica fotografabile"),
    ("Figar Café", "7. Bezirk", "Kirchengasse 18, 1070 Wien", "reservierung@figar.net", "+43 1 8909947", "figar.net", "@figar_vienna", "Café / Brunch / Bar", "Indipendente", "✅ Sì", "Café e bar trendy nel 7. distretto, ottimo per brunch instagrammabili"),
    ("Nook Coffee & Vintage", "13. Bezirk", "Altgasse 12, 1130 Wien", "studio@nookvienna.com", "+43 677 62793396", "nookvienna.com", "@nookvienna", "Specialty Coffee / Vintage", "Indipendente", "✅ Sì", "Concept store specialty coffee e vintage, molto curato ed estetico"),
    ("The Truth Café", "7. Bezirk", "Kaiserstraße 37, 1070 Wien", "hello@thetruthvienna.at", "+43 681 10808081", "thetruthvienna.at", "@thetruthvienna", "Specialty Coffee / Patisserie", "Piccola catena", "⚠️ Medio", "Caffè speciality con torte artigianali e cheesecake famose"),
    ("Chez Fritz Patisserie", "9. Bezirk", "Servitengasse 13, 1090 Wien", "", "+43 676 5434028", "chezfritz.at", "@chezfritz.patisserie", "Patisserie Vegana", "Indipendente", "✅ Sì", "Patisserie vegana 100% in stile francese nel Servitenviertel"),
    ("Superfood Deli", "1./6./9. Bezirk", "Wipplingerstr. 3 / Mariahilfer Str. 45, Wien", "hello@superfooddeli.at", "+43 1 3993313", "superfooddeli.at", "@superfooddeli", "Healthy / Vegan / Bowls", "Piccola catena", "⚠️ Medio", "Acai bowls, smoothies e piatti sani, molto instagrammabile"),
    ("Heunisch & Erben", "3. Bezirk", "Landstraßer Hauptstraße 17, 1030 Wien", "erben@heunisch.at", "+43 1 2868563", "heunisch.at", "", "Wine Bar / Restaurant", "Indipendente", "✅ Sì", "Enoteca e ristorante con 100+ vini al calice nel 3. distretto"),
    ("MAST Weinbistro", "9. Bezirk", "Porzellangasse 53, 1090 Wien", "contact@mast.wine", "+43 1 9226679", "mast.wine", "@mastvienna", "Wine Bar / Bistro", "Indipendente", "✅ Sì", "Weinbistro molto curato nell'Alsergrund, vini naturali e cucina stagionale"),
    ("Café Kandl", "7. Bezirk", "Kandlgasse 12, 1070 Wien", "tisch@cafekandl.at", "+43 1 8908938", "cafekandl.at", "", "Wine Bar / Café", "Indipendente", "✅ Sì", "Una delle destinazioni wine più eccitanti di Vienna, atmosfera bohémien"),
    ("Mochi", "2. Bezirk", "Praterstraße 15, 1020 Wien", "welcome@mochi.at", "+43 1 9251380", "mochi.at", "", "Japanese / Asian Fusion", "Piccola catena", "⚠️ Medio", "Punto di riferimento millennial viennesi, sushi e classici giapponesi"),
    ("Cucina Itameshi", "2. Bezirk", "Praterstraße 70, 1020 Wien", "info@cucina-itameshi.at", "+43 1 2122575", "cucina-itameshi.at", "@cucina_itameshi_vienna", "Italian-Japanese Fusion", "Indipendente", "✅ Sì", "Concept innovativo italo-giapponese aperto dal team di Mochi nel 2024"),
    ("Rinkhy Delicatessen Bar", "7. Bezirk", "Zieglergasse 29, 1070 Wien", "info@rinkhy.com", "+43 1 9478798", "rinkhy.com", "", "Tapas / Seafood / Bar", "Indipendente", "✅ Sì", "Bar di delicatessen con ostriche, tapas e ambiente energico nel 7. distretto"),
    ("Bruder Küche & Bar", "6. Bezirk", "Windmühlgasse 20, 1060 Wien", "tisch@bruder.xyz", "+43 664 1351320", "bruder.xyz", "", "New Nordic / Natural Wine", "Indipendente", "✅ Sì", "Cucina fermentata e vini naturali, muro di vasi fermentati come arredo"),
    ("Spelunke", "2. Bezirk", "Taborstraße 1, 1020 Wien", "ahoi@spelunke.at", "+43 1 2124151", "spelunke.at", "@spelunke_vienna", "Restaurant / Bar", "Indipendente", "✅ Sì", "Haubenrestaurant e bar instagrammabile nella Taborstraße"),
    ("Loca Casual Fine Dining", "1. Bezirk", "Stubenbastei 10, 1010 Wien", "office@bettereatbetter.com", "+43 1 5121172", "bettereatbetter.com", "@loca.casualfinedining", "Modern Austrian / Fine Casual", "Indipendente", "✅ Sì", "Fine dining informale nel centro con cucina moderna, arredamento ricercato"),
    ("Rosebar Centrala", "20. Bezirk", "Rauscherstraße 5, 1200 Wien", "hello@centrala.at", "+43 664 1266056", "centrala.at", "@rosebar.centrala", "Modern European", "Indipendente", "✅ Sì", "Nuovo ristorante con influenze dell'Europa orientale, aperto nel 2024"),
    ("Weimarck", "4. Bezirk", "Rechte Wienzeile 15, 1040 Wien", "office@weimarck.at", "+43 1 5874709", "weimarck.at", "", "Modern Austrian / Organic", "Indipendente", "✅ Sì", "Ristorante stagionale biologico vicino al Naschmarkt"),
    ("Labstelle", "1. Bezirk", "Lugeck 6, 1010 Wien", "hello@labstelle.at", "+43 1 2362122", "labstelle.at", "@labstelle", "Modern Viennese", "Indipendente", "✅ Sì", "Cucina viennese senza fronzoli, selezione di vini naturali"),
    ("Miznon", "1. Bezirk", "Schulerstraße 4, 1010 Wien", "vienna@miznonrestaurant.com", "+43 1 5121053", "miznonrestaurant.com", "@miznonvienna", "Israeli / Street Food", "Piccola catena", "⚠️ Medio", "Pita israeliana gourmet nel centro, instagrammabilissimo"),
    ("NENI am Wasser", "2. Bezirk", "Obere Donaustraße 65, 1020 Wien", "reservierung@neniamwasser.at", "+43 1 4380038", "nenifood.com", "@neni.wasser", "Tel Aviv / Mediterranean Fusion", "Piccola catena", "⚠️ Medio", "Ristorante sul canale del Danubio, cucina fusione Tel Aviv-Mediterraneo"),
    ("NENI am Naschmarkt", "6. Bezirk", "Naschmarkt 510, 1060 Wien", "naschmarkt@nenifood.com", "+43 1 5852020", "nenifood.com", "@neni.naschmarkt", "Tel Aviv / Mediterranean Fusion", "Piccola catena", "⚠️ Medio", "Location sul Naschmarkt, fusione mediterranea-mediorientale"),
    ("Maschu Maschu", "7. Bezirk", "Neubaugasse 20, 1070 Wien", "info@maschu-maschu.at", "+43 1 9904713", "maschu-maschu.at", "", "Middle Eastern / Falafel", "Piccola catena", "⚠️ Medio", "Falafel e cucina mediorientale nel 7. distretto"),
    ("Momoya", "1. Bezirk", "Börsegasse 3, 1010 Wien", "info@momoya.at", "+43 1 5350392", "momoya.at", "@restaurant_momoya", "Japanese / Sushi", "Indipendente", "✅ Sì", "Ristorante giapponese nel centro, cucina raffinata e ambiente elegante"),
    ("Nihonbashi", "1. Bezirk", "Kärntner Straße 44, 1010 Wien", "reservation@nihonbashi.at", "+43 1 8907856", "nihonbashi.at", "@nihonbashi.wien", "Japanese Fine Dining", "Indipendente", "✅ Sì", "Ristorante giapponese fine dining sulla Kärntner Straße, molto elegante"),
    ("Shiki Brasserie & Bar", "1. Bezirk", "Krugerstraße 3, 1010 Wien", "info@shiki.at", "+43 1 5127397", "shiki.at", "@shikivienna", "Japanese / Fine Dining", "Indipendente", "✅ Sì", "Cucina giapponese e brasserie nel centro, design moderno instagrammabile"),
    ("Ganko Yakiniku", "1. Bezirk", "Seilerstätte 2, 1010 Wien", "ganko.yakiniku@gmail.com", "+43 677 63096355", "", "@ganko_wien", "Japanese BBQ", "Indipendente", "✅ Sì", "Yakiniku (BBQ giapponese al tavolo) nel centro, unico nel suo genere"),
    ("Hollerei", "15. Bezirk", "Hollergasse 9, 1150 Wien", "info@hollerei.at", "+43 1 8923356", "hollerei.at", "@diehollerei", "Vegetarian / Art Gallery", "Indipendente", "✅ Sì", "Ristorante vegetariano e galleria d'arte dal 1999, concept unico"),
    ("Veggiezz", "1. Bezirk", "Salzgries 9A, 1010 Wien", "veggiezz.salzgries@outlook.com", "+43 1 5322650", "veggiezz.net", "", "Vegan / Bowls / Burgers", "Piccola catena", "⚠️ Medio", "Ristorante 100% vegano con bowls, burger e wrap"),
    ("Chez Bernard Rooftop", "5. Bezirk", "Hotel MOTTO, Schönbrunner Str. 30, 1050 Wien", "restaurant@chezbernard.at", "+43 1 5814600", "chezbernard.at", "", "Rooftop / Restaurant & Bar", "Indipendente", "✅ Sì", "Rooftop all'8° piano dell'Hotel MOTTO, una delle terrazze più instagrammabili"),
    ("Das Loft", "2. Bezirk", "Praterstraße 1 (SO/Vienna), 1020 Wien", "dasloft@so-hotels.com", "+43 664 88682556", "dasloftwien.at", "@dasloftwien", "Restaurant / Rooftop Bar", "Gruppo", "❌ No", "Bar-ristorante al 18° piano del SO/Vienna, parte di gruppo alberghiero"),
    ("Limón Rooftop", "1. Bezirk", "Schubertring 10-12 (Grand Ferdinand), 1010 Wien", "restaurant@cucinalimon.com", "+43 1 90211", "cucinalimon.com", "@cucinalimon", "Mediterranean / Rooftop", "Indipendente", "✅ Sì", "Nuovo rooftop mediterraneo al Grand Ferdinand, aperto 2024, molto fotografabile"),
    ("Glacis Beisl", "7. Bezirk", "Breite Gasse 4 (MuseumsQuartier), 1070 Wien", "mail@glacisbeisl.at", "+43 1 5265660", "glacisbeisl.at", "@glacis_beisl", "Viennese / Natural Wine", "Indipendente", "✅ Sì", "Uno dei giardini più belli di Vienna nel MuseumsQuartier, vini naturali"),
    ("Amerlingbeisl", "7. Bezirk", "Stiftgasse 8, 1070 Wien", "office@amerlingbeisl.at", "+43 1 5261660", "amerlingbeisl.at", "@amerlingbeisl", "Viennese Bistro / Garden", "Indipendente", "✅ Sì", "Bistrot viennese con cortile romantico nello Spittelberg, molto fotografabile"),
    ("Witwe Bolte", "7. Bezirk", "Gutenberggasse 13, 1070 Wien", "info@witwebolte.at", "+43 1 5231450", "witwebolte.at", "@witwebolte1070", "Traditional Austrian", "Indipendente", "✅ Sì", "Gasthaus tradizionale nello Spittelberg, atmosfera accogliente"),
    ("Bohème", "7. Bezirk", "Spittelberggasse 19, 1070 Wien", "tisch@boheme.at", "+43 699 16141922", "boheme.at", "", "Viennese / Austrian", "Indipendente", "✅ Sì", "Ristorante sulla pittoresca Spittelberggasse"),
    ("Gmoakeller", "3. Bezirk", "Am Heumarkt 25, 1030 Wien", "reservierung@gmoakeller.at", "+43 1 7125310", "gmoakeller.at", "@gmoakeller_wien", "Traditional Viennese", "Indipendente", "✅ Sì", "Wirtshaus viennese dal 1858, atmosfera autentica e molto fotografabile"),
    ("Gasthaus Pöschl", "1. Bezirk", "Weihburggasse 17, 1010 Wien", "andrea.poeschl@aon.at", "+43 1 5135288", "gaststättenpöschl.com", "@gasthaus.poeschl", "Traditional Austrian", "Indipendente", "✅ Sì", "Gasthaus tradizionale nel centro, cucina austriaca classica"),
    ("Huth Gastwirtschaft", "1. Bezirk", "Schellinggasse 5, 1010 Wien", "tisch@zum-huth.at", "+43 1 5135644", "zum-huth.at", "@huth_gastwirtschaft", "Austrian Gasthaus", "Indipendente", "✅ Sì", "Gastwirtschaft come dai libri di fiabe, cucina austriaca autentica"),
    ("Plachutta Wollzeile", "1. Bezirk", "Wollzeile 38, 1010 Wien", "wollzeile@plachutta.at", "+43 1 5121577", "plachutta-wollzeile.at", "@plachuttarestaurants", "Traditional Viennese / Tafelspitz", "Piccola catena", "⚠️ Medio", "Il tempio del Tafelspitz a Vienna"),
    ("Restaurant Wetter", "16. Bezirk", "Yppenplatz 11, 1160 Wien", "info@restaurantwetter.net", "+43 1 4060775", "restaurantwetter.net", "@restaurantwetterwien", "Mediterranean Italian", "Indipendente", "✅ Sì", "Ristorante mediterraneo-italiano sull'Yppenplatz nel quartiere hipster"),
    ("Café Hawelka", "1. Bezirk", "Dorotheergasse 6, 1010 Wien", "office@hawelka.at", "+43 1 5128230", "hawelka.at", "@cafehawelka", "Traditional Viennese Café", "Indipendente", "⚠️ Medio", "Storico Kaffeehaus dal 1939, frequentato da artisti, interior d'epoca"),
    ("Café Sperl", "6. Bezirk", "Gumpendorferstraße 11, 1060 Wien", "melange@cafesperl.at", "+43 1 5864158", "cafesperl.at", "@cafesperl", "Traditional Viennese Café", "Indipendente", "✅ Sì", "Uno dei più bei Kaffeehäuser di Vienna dal 1880, splendidamente fotografabile"),
    ("Café Schwarzenberg", "1. Bezirk", "Kärntner Ring 17, 1010 Wien", "office@cafe-schwarzenberg.at", "+43 1 5128998", "cafe-schwarzenberg.at", "@cafeschwarzenberg", "Grand Viennese Café", "Indipendente", "✅ Sì", "Elegante caffè sulla Ringstraße con interni storici instagrammabili"),
    ("Café Diglas", "1. Bezirk", "Wollzeile 10, 1010 Wien", "office@diglas.at", "+43 1 5125765", "diglas.at", "@cafe_diglas", "Traditional Viennese Café", "Piccola catena", "⚠️ Medio", "Caffè dal 1923 con sedie in pelle rossa e decorazioni vintage"),
    ("Café Landtmann", "1. Bezirk", "Universitätsring 4, 1010 Wien", "reservierung@landtmann.at", "+43 1 24100120", "landtmann.at", "@cafelandtmann", "Grand Viennese Café", "Indipendente", "⚠️ Medio", "Il più sofisticato Kaffeehaus di Vienna, vicino al Burgtheater"),
    ("Palmenhaus Brasserie", "1. Bezirk", "Burggarten, 1010 Wien", "office@palmenhaus.at", "+43 1 5331033", "palmenhaus.at", "@palmenhausbrasserie", "Café / Brasserie / Bar", "Indipendente", "✅ Sì", "Brasserie in una splendida serra storica nel Burggarten, molto instagrammabile"),
    ("Café Korb", "1. Bezirk", "Brandstätte 9, 1010 Wien", "cafe@cafekorb.at", "+43 1 5337215", "cafekorb.at", "@cafekorb", "Traditional Viennese Café", "Indipendente", "✅ Sì", "Kaffeehaus storico nel centro con atmosfera artistica"),
    ("Café Engländer", "1. Bezirk", "Postgasse 2, 1010 Wien", "postgasse@cafe-englaender.com", "+43 1 9668665", "cafe-englaender.com", "@cafe_englaender", "Viennese Café", "Piccola catena", "⚠️ Medio", "Moderno Kaffeehaus nel centro storico, menù classico e raffinato"),
    ("Café Frauenhuber", "1. Bezirk", "Himmelpfortgasse 6, 1010 Wien", "office@cafefrauenhuber.at", "+43 1 5125353", "cafefrauenhuber.at", "@cafefrauenhuber", "Traditional Viennese Café", "Indipendente", "✅ Sì", "Il più antico caffè di Vienna (1824), interni d'epoca molto fotografabili"),
    ("Café Rüdigerhof", "5. Bezirk", "Hamburger Straße 20, 1050 Wien", "office@caferuedigerhof.at", "+43 1 5863138", "caferüdigerhof.com", "@caferuedigerhof", "Traditional Viennese Café", "Indipendente", "✅ Sì", "Storico caffè con terrazza estiva, atmosfera Biedermeier instagrammabile"),
    ("Café Wortner", "4. Bezirk", "Wiedner Hauptstraße 55, 1040 Wien", "office@wortner.at", "+43 1 9458683", "wortner.at", "@cafewortner", "Traditional Viennese Café", "Indipendente", "✅ Sì", "Kaffeehaus storico, interni originali Art Nouveau instagrammabili"),
    ("Motto am Fluss", "1. Bezirk", "Franz-Josefs-Kai 2, 1010 Wien", "office@mottoamfluss.at", "+43 1 2525510", "mottoamfluss.at", "@mottoamfluss", "Restaurant / Café / Bar", "Indipendente", "✅ Sì", "Caffè e ristorante su una barca ancorata sul Danubio, originalissimo"),
    ("Café 7Stern", "7. Bezirk", "Siebensterngasse 31, 1070 Wien", "kulturcafe@7stern.net", "+43 699 15236157", "7stern.net", "@cafe7stern", "Café / Cultural Center", "Indipendente", "✅ Sì", "Centro culturale e caffè con giardino, eventi e atmosfera alternativa"),
    ("Café Halle MQ", "7. Bezirk", "Museumsplatz 1, 1070 Wien", "office@diehalle.at", "+43 1 5225327", "diehalle.at", "@halleeg_im_mq", "Café / Restaurant", "Indipendente", "✅ Sì", "Caffè nel MuseumsQuartier, ottima terrazza estiva fotografabile"),
    ("Café Leopold MQ", "7. Bezirk", "Museumsplatz 1, 1070 Wien", "reservierung@cafeleopold.wien", "+43 1 5222391", "cafeleopold.wien", "@cafe_leopold", "Café / Restaurant / Bar", "Indipendente", "✅ Sì", "Mix di caffè, ristorante e bar nel Leopold Museum"),
    ("Lingenhel", "3. Bezirk", "Landstraßer Hauptstraße 74, 1030 Wien", "office@lingenhel.com", "+43 1 7101566", "lingenhel.com", "@j.lingenhel", "Delicatessen / Wine Bar / Restaurant", "Indipendente", "✅ Sì", "Ristorante, vineria, gastronomia e primo caseificio urbano di Vienna, concept unico"),
    ("Zanoni & Zanoni", "1. Bezirk", "Am Lugeck 7, 1010 Wien", "garda.zanoni@chello.at", "+43 1 5127979", "zanoni.co.at", "@zanonieis", "Gelateria", "Piccola catena", "⚠️ Medio", "Storica gelateria artigianale italiana nel cuore di Vienna"),
    ("TIAN Bistro am Spittelberg", "7. Bezirk", "Schrankgasse 4, 1070 Wien", "spittelberg@taste-tian.com", "+43 1 5269491", "tian-bistro.com", "@tian_bistros", "Vegetarian Bistro", "Piccola catena", "⚠️ Medio", "Versione bistrot del Michelin-starred Tian, nel quartiere hipster di Spittelberg"),
    ("Freiraum", "6. Bezirk", "Mariahilfer Straße 117, 1060 Wien", "info@freiraum117.at", "+43 1 5969600", "freiraum117.at", "@freiraumwien", "Restaurant / Café / Bar", "Indipendente", "✅ Sì", "Ristorante trendy sulla Mariahilfer Straße con torrefazione propria"),
    ("Swing Kitchen", "8. Bezirk", "Josefstädter Str. 73, 1080 Wien", "office@swingkitchen.com", "", "swingkitchen.com", "@swing_kitchen", "Vegan Burgers", "Piccola catena", "⚠️ Medio", "Catena burger 100% vegani con 5 location a Vienna, molto instagrammabile"),
]


def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ristoranti Vienna - SocialPerks"

    # Colori
    HEADER_FILL   = PatternFill("solid", fgColor="2C3E50")
    EMAIL_FILL    = PatternFill("solid", fgColor="E8F5E9")
    NO_EMAIL_FILL = PatternFill("solid", fgColor="FFF9C4")
    IDEAL_FILL    = PatternFill("solid", fgColor="C8E6C9")
    MEDIUM_FILL   = PatternFill("solid", fgColor="FFF9C4")
    NO_FILL_CLR   = PatternFill("solid", fgColor="FFCDD2")
    WHITE_FONT    = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    BOLD_FONT     = Font(name="Calibri", bold=True, size=10)
    thin          = Side(style="thin", color="CCCCCC")
    BORDER        = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["#", "Nome Ristorante", "Distretto", "Indirizzo", "EMAIL ✉️",
               "Telefono", "Sito Web", "Instagram", "Tipo Cucina",
               "Dimensione", "Adatto SocialPerks", "Note/Descrizione"]
    col_widths = [4, 28, 14, 35, 32, 18, 22, 22, 22, 16, 14, 40]

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill   = HEADER_FILL
        cell.font   = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for idx, r in enumerate(RESTAURANTS, 1):
        row_idx = idx + 1
        has_email = bool(r[3])
        sp = r[9]
        row_fill = EMAIL_FILL if has_email else NO_EMAIL_FILL
        if "✅" in sp:
            sp_fill = IDEAL_FILL
        elif "⚠️" in sp:
            sp_fill = MEDIUM_FILL
        else:
            sp_fill = NO_FILL_CLR

        values = [idx, r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9], r[10]]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill   = sp_fill if col == 11 else row_fill
            cell.font   = Font(name="Calibri", size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col in [4, 12]))
        ws.row_dimensions[row_idx].height = 22

    total      = len(RESTAURANTS)
    with_email = sum(1 for r in RESTAURANTS if r[3])
    ideal      = sum(1 for r in RESTAURANTS if r[9].startswith("✅"))

    sr = total + 4
    ws.cell(row=sr,   column=1, value="📊 RIEPILOGO").font = Font(bold=True, size=12, color="2C3E50")
    ws.cell(row=sr+1, column=1, value="Totale ristoranti:").font     = BOLD_FONT
    ws.cell(row=sr+1, column=2, value=total)
    ws.cell(row=sr+2, column=1, value="Con email confermata:").font  = BOLD_FONT
    ws.cell(row=sr+2, column=2, value=with_email)
    ws.cell(row=sr+3, column=1, value="Ideali per SocialPerks (✅):").font = BOLD_FONT
    ws.cell(row=sr+3, column=2, value=ideal)
    ws.cell(row=sr+4, column=1, value="Generato il:").font           = BOLD_FONT
    ws.cell(row=sr+4, column=2, value=datetime.now().strftime("%d/%m/%Y %H:%M"))

    filename = "/home/user/claude-restaurants-/SocialPerks_Restaurants_Vienna.xlsx"
    wb.save(filename)
    print(f"✅ {filename}")
    print(f"   Totale: {total} | Con email: {with_email} | Ideali: {ideal}")
    return filename


if __name__ == "__main__":
    create_excel()
